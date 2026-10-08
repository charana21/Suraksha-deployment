import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.forecasting.data_loader import _get_sync_db

def create_excel_report(output_file: str):
    db = _get_sync_db()

    # 1. Fetch 2026 Forecasting Data
    f_docs = list(db.forecasting_data.find({'date': {'$regex': '^2026-'}}).sort('date', 1))
    df_forecast = pd.DataFrame(f_docs)

    # 2. Fetch Historical Data (2022 to latest)
    h_docs = list(db.historical_data.find({}).sort('date', 1))
    df_hist = pd.DataFrame(h_docs)

    # Prepare Sheet 1 Data
    rows_sheet1 = []
    for doc in f_docs:
        d = doc.get('date')
        dow = doc.get('day_of_week')
        occ = doc.get('occasion_or_holiday')
        if pd.isna(occ) or str(occ).strip().lower() in ['nan', 'none', '']:
            occ = ''
        
        uts_act = doc.get('uts_total_actual')
        uts_pred = doc.get('uts_total_predicted')
        prs_act = doc.get('prs_total_actual')
        prs_pred = doc.get('prs_total_predicted')
        ovr_act = doc.get('overall_total_actual')
        ovr_pred = doc.get('overall_total_predicted')
        is_blind = doc.get('is_blind_forecast', False)

        uts_diff = (uts_pred - uts_act) if (uts_act is not None and not pd.isna(uts_act)) else None
        prs_diff = (prs_pred - prs_act) if (prs_act is not None and not pd.isna(prs_act)) else None
        ovr_diff = (ovr_pred - ovr_act) if (ovr_act is not None and not pd.isna(ovr_act)) else None

        uts_pct = (abs(uts_diff) / uts_act * 100.0) if (uts_diff is not None and uts_act and uts_act > 0) else None
        prs_pct = (abs(prs_diff) / prs_act * 100.0) if (prs_diff is not None and prs_act and prs_act > 0) else None
        ovr_pct = (abs(ovr_diff) / ovr_act * 100.0) if (ovr_diff is not None and ovr_act and ovr_act > 0) else None

        status = "Blind Future Forecast" if is_blind else "Historical Period (Actual Available)"
        
        if ovr_pred is not None:
            if ovr_pred > 160000:
                alert = "Red (>160k Emergency)"
            elif ovr_pred >= 140000:
                alert = "Orange (140k-160k Alert)"
            else:
                alert = "Green (<140k Normal)"
        else:
            alert = ""

        rows_sheet1.append({
            'Date': d,
            'Day of Week': dow,
            'Occasion / Holiday': occ,
            'Forecast Type': status,
            'UTS Actual (Original)': uts_act,
            'UTS Predicted': uts_pred,
            'UTS Difference': uts_diff,
            'UTS Error %': round(uts_pct, 2) if uts_pct is not None else None,
            'PRS Actual (Original)': prs_act,
            'PRS Predicted': prs_pred,
            'PRS Difference': prs_diff,
            'PRS Error %': round(prs_pct, 2) if prs_pct is not None else None,
            'Overall Actual (Original)': ovr_act,
            'Overall Predicted': ovr_pred,
            'Overall Difference': ovr_diff,
            'Overall Error %': round(ovr_pct, 2) if ovr_pct is not None else None,
            'Day Alert Status': alert
        })

    df_sheet1 = pd.DataFrame(rows_sheet1)

    # Prepare Sheet 2 Data (Monthly Summary)
    df_sheet1['Month_Num'] = pd.to_datetime(df_sheet1['Date']).dt.month
    month_names = ['January', 'February', 'March', 'April', 'May', 'June', 
                   'July', 'August', 'September', 'October', 'November', 'December']
    
    rows_sheet2 = []
    for m in range(1, 13):
        m_df = df_sheet1[df_sheet1['Month_Num'] == m]
        days = len(m_df)
        
        u_act_sum = m_df['UTS Actual (Original)'].dropna().sum() if m_df['UTS Actual (Original)'].dropna().count() > 0 else None
        u_pred_sum = m_df['UTS Predicted'].sum()
        p_act_sum = m_df['PRS Actual (Original)'].dropna().sum() if m_df['PRS Actual (Original)'].dropna().count() > 0 else None
        p_pred_sum = m_df['PRS Predicted'].sum()
        o_act_sum = m_df['Overall Actual (Original)'].dropna().sum() if m_df['Overall Actual (Original)'].dropna().count() > 0 else None
        o_pred_sum = m_df['Overall Predicted'].sum()
        o_pred_avg = m_df['Overall Predicted'].mean()

        if o_pred_avg >= 145000:
            m_badge = "Red (High Demand >=145k/day)"
        elif o_pred_avg >= 140000:
            m_badge = "Orange (Moderate Demand 140-145k/day)"
        else:
            m_badge = "Green (Normal Demand <140k/day)"

        rows_sheet2.append({
            'Month': f"{m:02d} - {month_names[m-1]} 2026",
            'Total Days': days,
            'UTS Actual Total': u_act_sum,
            'UTS Predicted Total': round(u_pred_sum),
            'PRS Actual Total': p_act_sum,
            'PRS Predicted Total': round(p_pred_sum),
            'Overall Actual Total': o_act_sum,
            'Overall Predicted Total': round(o_pred_sum),
            'Predicted Daily Average': round(o_pred_avg),
            'Monthly Badge': m_badge
        })

    df_sheet2 = pd.DataFrame(rows_sheet2)

    # Prepare Sheet 3 Data (Complete Historical Actuals)
    rows_sheet3 = []
    for doc in h_docs:
        occ = doc.get('occasion_or_holiday')
        if pd.isna(occ) or str(occ).strip().lower() in ['nan', 'none', '']:
            occ = ''
        rows_sheet3.append({
            'Date': doc.get('date'),
            'Day of Week': doc.get('day_of_week'),
            'UTS Total (Original Actual)': doc.get('uts_total'),
            'PRS Total (Original Actual)': doc.get('prs_total'),
            'Overall Total (Original Actual)': doc.get('overall_total'),
            'Occasion / Holiday': occ
        })
    df_sheet3 = pd.DataFrame(rows_sheet3)

    # Create Excel Workbook using openpyxl
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # STYLES
    navy_header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    dark_teal_fill = PatternFill(start_color="134F5C", end_color="134F5C", fill_type="solid")
    steel_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    green_soft = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    orange_soft = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    red_soft = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1F4E79"))

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    def format_sheet(ws, df, header_fill, sheet_title):
        ws.title = sheet_title
        ws.views.sheetView[0].showGridLines = True

        # Write Headers
        headers = list(df.columns)
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = cell_border
        ws.row_dimensions[1].height = 28

        # Write Data
        for r_idx, row in enumerate(df.values, start=2):
            ws.append(list(row))
            ws.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            
            for c_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = regular_font
                cell.border = cell_border

                val = row[c_idx - 1]
                col_name = headers[c_idx - 1]

                # Alignments and Number Formats
                if isinstance(val, (int, np.integer)):
                    cell.alignment = align_right
                    cell.number_format = "#,##0"
                elif isinstance(val, (float, np.floating)):
                    cell.alignment = align_right
                    if "%" in col_name or "Error %" in col_name:
                        cell.number_format = "0.00\"%\""
                    else:
                        cell.number_format = "#,##0"
                elif col_name in ['Date', 'Day of Week', 'Forecast Type']:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

                # Zebra striping
                if is_zebra and not cell.fill.start_color.rgb:
                    cell.fill = zebra_fill

                # Special Highlights
                if col_name == 'Day Alert Status':
                    if "Green" in str(val):
                        cell.fill = green_soft
                    elif "Orange" in str(val):
                        cell.fill = orange_soft
                    elif "Red" in str(val):
                        cell.fill = red_soft
                elif col_name == 'Monthly Badge':
                    if "Green" in str(val):
                        cell.fill = green_soft
                    elif "Orange" in str(val):
                        cell.fill = orange_soft
                    elif "Red" in str(val):
                        cell.fill = red_soft

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Freeze top row
        ws.freeze_panes = "A2"

    # 1. Sheet 1: 2026 Forecast vs Actuals
    ws1 = wb.create_sheet(title="2026_Forecast_vs_Actual")
    format_sheet(ws1, df_sheet1.drop(columns=['Month_Num']), navy_header_fill, "2026_Forecast_vs_Actual")

    # 2. Sheet 2: Monthly Summary
    ws2 = wb.create_sheet(title="2026_Monthly_Summary")
    format_sheet(ws2, df_sheet2, dark_teal_fill, "2026_Monthly_Summary")

    # 3. Sheet 3: Historical Actuals
    ws3 = wb.create_sheet(title="Historical_Actuals_2022_2026")
    format_sheet(ws3, df_sheet3, steel_fill, "Historical_Actuals_2022_2026")

    wb.save(output_file)
    print(f"Excel successfully created at: {output_file}")

if __name__ == '__main__':
    out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "UTS_PRS_Original_vs_Predicted_2026.xlsx")
    create_excel_report(out_path)
